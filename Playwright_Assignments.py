from playwright.sync_api import sync_playwright
from openpyxl import load_workbook, Workbook
from datetime import datetime
import json
import os
import time


# -----------------------------
# Read contacts.xlsx
# -----------------------------

workbook = load_workbook("contacts.xlsx")
sheet = workbook.active

contacts = []

for row in sheet.iter_rows(min_row=2, values_only=True):

    name, phone, message = row

    if not name or not phone:
        continue

    contacts.append({
        "name": str(name),
        "phone": str(phone),
        "message": str(message or "Hello {name}")
    })


# -----------------------------
# Create screenshot folder
# -----------------------------

os.makedirs("screenshots", exist_ok=True)

report = []


# -----------------------------
# Open WhatsApp Web
# -----------------------------

with sync_playwright() as p:

    context = p.chromium.launch_persistent_context(
        user_data_dir="whatsapp_profile",
        channel="msedge",
        headless=False
    )

    page = context.pages[0] if context.pages else context.new_page()

    page.goto("https://web.whatsapp.com/")

    print("Please scan the QR code manually if required...")

    page.wait_for_selector("#side", timeout=120000)

    print("WhatsApp Web login successful!")


    # -----------------------------
    # Process each contact
    # -----------------------------

    for contact in contacts:

        name = contact["name"]
        phone = contact["phone"]

        message = contact["message"].replace(
            "{name}",
            name
        )

        print("\nProcessing:", name)

        try:

            # -----------------------------
            # Find WhatsApp search box
            # -----------------------------

            print("Searching for:", phone)

            search_box = page.get_by_role(
                "textbox",
                name="Search"
            )

            search_box.wait_for(
                state="visible",
                timeout=15000
            )

            search_box.click()

            search_box.fill(phone)

            time.sleep(3)


            # -----------------------------
            # Open search result
            # -----------------------------

            result = page.get_by_text(
                name,
                exact=False
            ).first

            if result.count() > 0:

                result.click()

            else:

                result = page.get_by_text(
                    phone,
                    exact=False
                ).first

                result.click()

            time.sleep(3)


            # -----------------------------
            # Find message box
            # -----------------------------

            message_box = page.locator(
                'footer div[contenteditable="true"]'
            ).last

            message_box.wait_for(
                state="visible",
                timeout=15000
            )

            message_box.click()

            message_box.fill(message)

            message_box.press("Enter")

            print("Message sent to:", name)

            time.sleep(3)


            # -----------------------------
            # Take screenshot
            # -----------------------------

            safe_name = "".join(
                c if c.isalnum() else "_"
                for c in name
            )

            screenshot_file = (
                f"screenshots/"
                f"{safe_name}_"
                f"{datetime.now().strftime('%Y%m%d_%H%M%S')}.png"
            )

            page.screenshot(
                path=screenshot_file,
                full_page=False
            )

            print(
                "Screenshot saved:",
                screenshot_file
            )


            # -----------------------------
            # Extract last 3 incoming messages
            # -----------------------------

            incoming = page.locator(
                "div.message-in"
            )

            incoming_messages = []

            count = incoming.count()

            start = max(0, count - 3)

            for i in range(start, count):

                try:

                    text = incoming.nth(i).inner_text().strip()

                    if text:
                        incoming_messages.append(text)

                except Exception:
                    pass


            # -----------------------------
            # Add result to report
            # -----------------------------

            report.append({

                "name": name,

                "phone": phone,

                "message_sent": message,

                "last_3_messages": incoming_messages,

                "screenshot": screenshot_file,

                "timestamp":
                    datetime.now().strftime(
                        "%Y-%m-%d %H:%M:%S"
                    ),

                "status": "Sent"

            })


        except Exception as error:

            print(
                "Error for",
                name,
                ":",
                error
            )

            report.append({

                "name": name,

                "phone": phone,

                "message_sent": message,

                "last_3_messages": [],

                "screenshot": "",

                "timestamp":
                    datetime.now().strftime(
                        "%Y-%m-%d %H:%M:%S"
                    ),

                "status": f"Failed: {error}"

            })


    # -----------------------------
    # Save JSON report
    # -----------------------------

    with open(
        "whatsapp_report.json",
        "w",
        encoding="utf-8"
    ) as file:

        json.dump(
            report,
            file,
            indent=4,
            ensure_ascii=False
        )


    # -----------------------------
    # Save Excel report
    # -----------------------------

    report_workbook = Workbook()

    report_sheet = report_workbook.active

    report_sheet.title = "WhatsApp Report"

    report_sheet.append([
        "Name",
        "Phone",
        "Message Sent",
        "Last Message 1",
        "Last Message 2",
        "Last Message 3",
        "Screenshot",
        "Timestamp",
        "Status"
    ])


    for item in report:

        messages = item["last_3_messages"].copy()

        while len(messages) < 3:
            messages.append("")

        report_sheet.append([
            item["name"],
            item["phone"],
            item["message_sent"],
            messages[0],
            messages[1],
            messages[2],
            item["screenshot"],
            item["timestamp"],
            item["status"]
        ])


    report_workbook.save(
        "whatsapp_report.xlsx"
    )


    # -----------------------------
    # Finished
    # -----------------------------

    print("\nAssignment completed!")

    print(
        "JSON report: whatsapp_report.json"
    )

    print(
        "Excel report: whatsapp_report.xlsx"
    )

    print(
        "Screenshots: screenshots folder"
    )

    input(
        "Press Enter to close..."
    )

    context.close()